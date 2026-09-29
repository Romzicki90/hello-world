"""Reads SAP list-display / table exports and normalises them to a standard dataset.

Handles what SAP exports actually look like:
  * .xlsx from the ALV "Spreadsheet" export, with title lines above the header;
  * .csv / .tsv / .txt, including the "unconverted" list format with | separators
    and ----- rule lines;
  * repeated header lines, subtotal / total lines, blank lines;
  * SAP decimal notations (1,234,567.89 / 1.234.567,89 / 1 234 567,89), Indian digit
    grouping, trailing minus (1,234.00-), brackets for negatives;
  * SAP date formats (DD.MM.YYYY etc.), SE16N internal dates (YYYYMMDD),
    empty dates (00.00.0000), Excel date cells;
  * leading zeros on numeric identifiers (vendor 0000100005 = 100005).

Every dropped line and every value that could not be read is counted and reported
in the Data Quality sheet; nothing is silently discarded.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from .loader import Dataset, normalise_header

DECIMAL_NOTATIONS = {
    "1,234,567.89": "X",  # SAP user setting X (usual in India)
    "1.234.567,89": " ",  # SAP user setting blank
    "1 234 567,89": "Y",  # SAP user setting Y
}

DATE_FORMATS = {
    "DD.MM.YYYY": "%d.%m.%Y",
    "MM/DD/YYYY": "%m/%d/%Y",
    "MM-DD-YYYY": "%m-%d-%Y",
    "YYYY.MM.DD": "%Y.%m.%d",
    "YYYY/MM/DD": "%Y/%m/%d",
    "YYYY-MM-DD": "%Y-%m-%d",
    "DD/MM/YYYY": "%d/%m/%Y",
    "DD-MM-YYYY": "%d-%m-%Y",
}

CREDIT_INDICATORS = {"H", "C", "CR", "CREDIT"}
EMPTY_DATES = {"", "0", "00000000", "00.00.0000", "00/00/0000", "0000-00-00", "#"}
HEADER_SCAN_ROWS = 40


class IngestError(Exception):
    pass


@dataclass
class FileReport:
    path: str
    sha256: str
    sheet: str | None
    header_row: int
    rows_read: int = 0
    rows_loaded: int = 0
    dropped: Counter = field(default_factory=Counter)
    mapping: dict[str, str] = field(default_factory=dict)
    unmapped_columns: list[str] = field(default_factory=list)


@dataclass
class DatasetLoad:
    dataset: Dataset
    files: list[FileReport]
    rows: list[dict]
    parse_failures: dict[str, Counter]
    mapped_fields: set[str]
    warnings: list[str]


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------------------- reading

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return raw.decode("cp1252", errors="replace")  # SAP GUI default for Windows exports


def _is_rule_line(line: str) -> bool:
    s = line.strip()
    return bool(s) and set(s) <= set("-=_|+ ")


def read_rows(path: Path, sheet: str | None = None) -> list[list]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    if suffix == ".xls":
        raise IngestError(f"{path.name}: old .xls format. Open it in Excel and save as .xlsx.")
    if suffix not in (".csv", ".tsv", ".txt", ".dat"):
        raise IngestError(f"{path.name}: unsupported file type {suffix}")

    text = _read_text(path)
    lines = text.splitlines()
    sample = [ln for ln in lines[:200] if ln.strip() and not _is_rule_line(ln)]
    counts = {d: sum(ln.count(d) for ln in sample) for d in ("|", "\t", ";", ",")}
    if counts["|"] >= max(len(sample), 1):
        rows = []
        for ln in lines:
            if not ln.strip() or _is_rule_line(ln):
                rows.append([])
                continue
            s = ln.strip()
            if s.startswith("|"):
                s = s[1:]
            if s.endswith("|"):
                s = s[:-1]
            rows.append([c.strip() for c in s.split("|")])
        return rows
    delim = "\t" if counts["\t"] >= max(counts[";"], counts[","]) and counts["\t"] else (
        ";" if counts[";"] > counts[","] else ",")
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


# --------------------------------------------------------------------------- parsing

def parse_number(value, notation: str):
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("boolean")
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", " ")
    if s in ("", "-"):
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    if s.endswith("-"):
        negative, s = True, s[:-1].strip()
    if s.startswith("-"):
        negative, s = not negative, s[1:].strip()
    s = re.sub(r"(?i)^(rs\.?|inr|₹)\s*", "", s)
    s = re.sub(r"(?i)\s*(rs\.?|inr|₹)$", "", s)
    code = DECIMAL_NOTATIONS.get(notation, notation)
    if code == "X":
        s = s.replace(",", "").replace(" ", "")
    elif code == " ":
        s = s.replace(".", "").replace(" ", "").replace(",", ".")
    elif code == "Y":
        s = s.replace(" ", "").replace(",", ".")
    number = float(s)  # raises ValueError on anything left over
    return -number if negative else number


def parse_date(value, date_format: str):
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        n = int(value)
        if 19000101 <= n <= 99991231:
            return dt.datetime.strptime(str(n), "%Y%m%d").date()
        if 1 <= n < 2958466:  # Excel serial date
            return dt.date(1899, 12, 30) + dt.timedelta(days=n)
        raise ValueError(f"number {value}")
    s = str(value).strip()
    if s in EMPTY_DATES:
        return None
    s = s.split(" ")[0]  # drop a time part
    for fmt in (DATE_FORMATS.get(date_format, date_format), "%Y%m%d", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(s)


def normalise_id(value, strip_prefixes=(), keep_zeros=False):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    for prefix in strip_prefixes:
        if s.upper().startswith(prefix.upper()):
            s = s[len(prefix):].strip()
    if not s:
        return None
    s = re.sub(r"\s+", " ", s.upper())
    if s.isdigit() and not keep_zeros:
        s = s.lstrip("0") or "0"
    return s


def normalise_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    s = re.sub(r"\s+", " ", str(value)).strip()
    return s or None


# --------------------------------------------------------------------------- mapping

def _cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def detect_header(rows: list[list], dataset: Dataset, explicit: dict[str, str]) -> int:
    """Index of the row that best matches the dataset's known column names."""
    wanted = {normalise_header(s) for f in dataset.fields.values() for s in f.synonyms}
    wanted |= {normalise_header(c) for c in explicit.values()}
    best, best_score = -1, 0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = len({normalise_header(c) for c in row if _cell_text(c)} & wanted)
        if score > best_score:
            best, best_score = i, score
    if best_score < 2:
        raise IngestError(
            f"could not find a header row matching dataset '{dataset.id}' in the first "
            f"{HEADER_SCAN_ROWS} rows. Map the columns explicitly in the profile.")
    return best


def map_columns(header: list[str], dataset: Dataset, explicit: dict[str, str]) -> dict[str, int]:
    """Standard field name -> column index. Explicit mapping wins over synonyms."""
    norm = [normalise_header(h) for h in header]
    used: set[int] = set()
    result: dict[str, int] = {}
    for fname, column in explicit.items():
        if fname not in dataset.fields:
            raise IngestError(f"mapping refers to unknown field '{fname}' of dataset '{dataset.id}'")
        key = normalise_header(column)
        if key not in norm:
            raise IngestError(f"mapped column '{column}' for field '{fname}' not found in the header")
        idx = norm.index(key)
        result[fname] = idx
        used.add(idx)
    for fname, f in dataset.fields.items():
        if fname in result:
            continue
        for syn in [fname] + f.synonyms:
            key = normalise_header(syn)
            idx = next((i for i, n in enumerate(norm) if n == key and i not in used), None)
            if idx is not None:
                result[fname] = idx
                used.add(idx)
                break
    return result


# --------------------------------------------------------------------------- loading

def load_dataset(dataset: Dataset, files: list[Path], *, sheet: str | None = None,
                 explicit: dict[str, str] | None = None, decimal_notation: str = "1,234,567.89",
                 date_format: str = "DD.MM.YYYY") -> DatasetLoad:
    explicit = explicit or {}
    all_rows: list[dict] = []
    reports: list[FileReport] = []
    failures: dict[str, Counter] = defaultdict(Counter)
    mapped: set[str] | None = None
    key_fields = [n for n, f in dataset.fields.items() if f.key]

    for path in files:
        raw = read_rows(path, sheet)
        h = detect_header(raw, dataset, explicit)
        header = [_cell_text(c) for c in raw[h]]
        colmap = map_columns(header, dataset, explicit)
        rep = FileReport(path=str(path), sha256=sha256_of(path), sheet=sheet, header_row=h + 1,
                         mapping={f: header[i] for f, i in colmap.items()},
                         unmapped_columns=[c for i, c in enumerate(header) if c and i not in colmap.values()])
        reports.append(rep)
        norm_header = [normalise_header(c) for c in header]
        this_mapped = set(colmap)
        mapped = this_mapped if mapped is None else mapped & this_mapped

        for row in raw[h + 1:]:
            cells = [_cell_text(c) for c in row]
            if not any(cells):
                rep.dropped["blank line"] += 1
                continue
            rep.rows_read += 1
            if all(set(c) <= set("-=_*|+ ") for c in cells if c):
                rep.dropped["rule / separator line"] += 1
                continue
            nonblank = [(i, normalise_header(c)) for i, c in enumerate(cells) if c]
            if nonblank and sum(1 for i, n in nonblank if i < len(norm_header) and n == norm_header[i]) >= max(2, len(nonblank) * 0.6):
                rep.dropped["repeated header line"] += 1
                continue
            if any(c.startswith("*") for c in cells[:3] if c):
                rep.dropped["subtotal / total line (marked *)"] += 1
                continue

            rec = {}
            for fname, idx in colmap.items():
                f = dataset.fields[fname]
                value = row[idx] if idx < len(row) else None
                if isinstance(value, str):
                    value = value.strip()
                try:
                    if f.type == "id":
                        rec[fname] = normalise_id(value, f.strip_prefixes, f.keep_zeros)
                    elif f.type == "text":
                        rec[fname] = normalise_text(value)
                    elif f.type == "date":
                        rec[fname] = parse_date(value, date_format)
                    else:
                        rec[fname] = parse_number(value, decimal_notation)
                except (ValueError, TypeError):
                    rec[fname] = None
                    failures[fname][_cell_text(value)[:40]] += 1
            if key_fields and all(rec.get(k) is None for k in key_fields if k in colmap):
                rep.dropped["blank key fields (subtotal / total line)"] += 1
                continue
            for fname in colmap:
                sign_field = dataset.fields[fname].sign_by
                if sign_field and sign_field in colmap and rec.get(fname) and rec[fname] > 0:
                    if (rec.get(sign_field) or "").upper() in CREDIT_INDICATORS:
                        rec[fname] = -rec[fname]
            all_rows.append(rec)
            rep.rows_loaded += 1

    warnings = []
    for fname, f in dataset.fields.items():
        if f.sign_by and fname in (mapped or set()):
            values = [r[fname] for r in all_rows if r.get(fname) is not None]
            if values and min(values) >= 0 and f.sign_by not in (mapped or set()):
                warnings.append(
                    f"'{fname}' has no negative values and no debit/credit column is mapped: "
                    f"amounts look unsigned. Tests that rely on credit (negative) items will miss them.")
    return DatasetLoad(dataset=dataset, files=reports, rows=all_rows, parse_failures=dict(failures),
                       mapped_fields=mapped or set(), warnings=warnings)
