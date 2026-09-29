"""Every analytic test must find exactly the exceptions planted in the synthetic data:
all planted exceptions, and none of the near-misses or background transactions."""

import pytest

from auditwb.engine import EXCEPTIONS, NOT_EXECUTABLE, run
from auditwb.loader import load_library
from auditwb.profile import load_profile
from auditwb.report import annexure_rows, write_annexure_ii, write_working_paper
from auditwb.synthetic import PLANTED, generate


@pytest.fixture(scope="module")
def demo_run(tmp_path_factory):
    target = tmp_path_factory.mktemp("demo")
    profile = load_profile(generate(target))
    return run(profile, load_library(), progress=lambda *_: None)


def test_every_library_test_has_planted_cases():
    assert set(PLANTED) == set(load_library().tests)


@pytest.mark.parametrize("test_id", sorted(PLANTED))
def test_finds_exactly_the_planted_exceptions(demo_run, test_id):
    result = demo_run.result(test_id)
    assert result.status == EXCEPTIONS, result.message
    column, expected = PLANTED[test_id]
    found = {str(row[result.columns.index(column)]) for row in result.rows}
    assert found == expected


def test_population_and_value_are_reported(demo_run):
    for r in demo_run.results:
        assert r.population_n, f"{r.test.id}: population not reported"
        assert r.value is not None and r.value > 0, f"{r.test.id}: value involved missing"


def test_split_po_cluster_value(demo_run):
    r = demo_run.result("PRC-SPLIT-01")
    assert r.n_exceptions == 1
    assert r.value == pytest.approx(1_050_000)


def test_ingest_drops_subtotal_and_total_lines(demo_run):
    load = demo_run.datasets["vendor_items"].load
    dropped = sum(sum(f.dropped.values()) for f in load.files)
    assert dropped >= 2
    assert "vendor" in load.mapped_fields and "amount" in load.mapped_fields


def test_missing_data_is_not_executable_not_nil(tmp_path):
    profile = load_profile(generate(tmp_path))
    del profile.datasets["grir"]
    profile.params["FIN-SENSGL-01"] = {"sensitive_gls": []}
    result = run(profile, load_library(), progress=lambda *_: None)
    assert result.result("PRC-GRIR-01").status == NOT_EXECUTABLE
    assert "grir" in result.result("PRC-GRIR-01").message
    assert result.result("FIN-SENSGL-01").status == NOT_EXECUTABLE


def test_settings_saved_from_window_reload_identically(tmp_path):
    from auditwb.profile import dump_toml, read_raw

    raw = read_raw(generate(tmp_path / "demo"))
    raw["params"]["PRC-SPLIT-01"] = {"threshold": 2000000}
    saved = tmp_path / "saved" / "audit_settings.toml"
    saved.parent.mkdir()
    saved.write_text(dump_toml(raw), encoding="utf-8")
    assert read_raw(saved) == raw
    result = run(load_profile(saved), load_library(), progress=lambda *_: None)
    assert result.result("PRC-SPLIT-01").n_exceptions == 0   # 10.5 lakh cluster is below 20 lakh


def _demo(tmp_path):
    from auditwb.profile import read_raw

    path = generate(tmp_path / "demo")
    return path, read_raw(path)


def _run_raw(raw, path):
    from auditwb.profile import profile_from_dict

    return run(profile_from_dict(raw, path), load_library(), progress=lambda *_: None)


def test_demo_inputs_are_reconciled(demo_run):
    for st in demo_run.datasets.values():
        if st.load:
            assert st.recon_status in ("RECONCILED", "AUDIT CRITERIA"), (st.id, st.recon_note)
    assert all(r.reconciled for r in demo_run.results)


def test_overlapping_price_criteria_block_the_test(tmp_path):
    from openpyxl import load_workbook

    path, raw = _demo(tmp_path)
    xlsx = tmp_path / "demo" / "criteria" / "approved_prices.xlsx"
    wb = load_workbook(xlsx)
    wb.active.append(["000000000020000001", "TO", "01.10.2025", "31.03.2026", 48000, "Second circular"])
    wb.save(xlsx)
    result = _run_raw(raw, path)
    r = result.result("SAL-PRICE-01")
    assert r.status == "BLOCKED"
    assert "overlapping validity" in r.message
    assert result.result("SAL-CN-01").status == "EXCEPTIONS"      # other tests unaffected


def test_duplicate_master_keys_block_dependent_tests(tmp_path):
    from openpyxl import load_workbook

    path, raw = _demo(tmp_path)
    xlsx = tmp_path / "demo" / "extracts" / "KNA1.xlsx"
    wb = load_workbook(xlsx)
    wb.active.append(["300009", "Blocked Customer (duplicate row)", "01.04.2015", "X", "01", "", ""])
    wb.save(xlsx)
    raw["datasets"]["customer_master"]["control_rows"] += 1
    result = _run_raw(raw, path)
    assert result.result("SAL-BLOCKED-01").status == "BLOCKED"
    assert "duplicate key" in result.result("SAL-BLOCKED-01").message


def test_nil_requires_reconciliation(tmp_path):
    path, raw = _demo(tmp_path)
    raw["params"]["SAL-AR-01"] = {"overdue_days": 100000}
    r = _run_raw(raw, path).result("SAL-AR-01")
    assert r.display_status == "NIL - RECONCILED"

    del raw["datasets"]["customer_items"]["control_rows"]
    del raw["datasets"]["customer_items"]["control_total"]
    r = _run_raw(raw, path).result("SAL-AR-01")
    assert r.display_status == "NIL - NOT RECONCILED"

    raw["datasets"]["customer_items"]["control_rows"] = 61      # SAP shows one row more than loaded
    raw["datasets"]["customer_items"]["control_total"] = 0
    result = _run_raw(raw, path)
    assert result.datasets["customer_items"].recon_status == "DIFFERENCE"
    assert result.result("SAL-AR-01").display_status == "NIL - NOT RECONCILED"
    rows = {x["topic"].id: x for x in annexure_rows(result)}
    assert "not analytical assurance" in rows["T3A08"]["detail"]


def test_return_period_needs_the_policy_criterion(tmp_path):
    path, raw = _demo(tmp_path)
    del raw["params"]["SAL-RETURN-01"]
    r = _run_raw(raw, path).result("SAL-RETURN-01")
    assert r.status == NOT_EXECUTABLE and "max_return_days" in r.message


def test_unmapped_column_disables_only_its_check(tmp_path):
    path, raw = _demo(tmp_path)
    raw["datasets"]["sales_returns"]["columns"] = {"qc_status": ""}   # 'not in file' chosen in the window
    r = _run_raw(raw, path).result("SAL-RETURN-02")
    found = {row[r.columns.index("exception_id")] for row in r.rows}
    assert "7100000006/10" not in found            # the QC check no longer runs
    assert "7100000005/10" in found                # the other checks still do


def test_self_test_passes():
    from auditwb.synthetic import self_test

    ok, lines = self_test()
    assert ok, "\n".join(lines)


def test_cmo_starter_profile_loads_and_marks_return_period_as_required(tmp_path):
    from pathlib import Path

    from auditwb.profile import read_raw

    root = next(p for p in Path(__file__).parents if (p / "profiles").exists())   # repo or release layout
    path = root / "profiles" / "SAIL_CMO_2025-26.toml"
    raw = read_raw(path)
    assert raw["annexure_ii"]["topics"] == ["T3A01", "T3A04", "T3A05", "T3A08"]
    assert "max_return_days" not in raw["params"]["SAL-RETURN-01"]


def test_outputs_are_written(demo_run, tmp_path):
    wp = write_working_paper(demo_run, tmp_path)
    ax = write_annexure_ii(demo_run, tmp_path)
    assert wp.exists() and ax.exists()
    rows = {r["topic"].id: r for r in annexure_rows(demo_run)}
    assert rows["T1A01"]["exceptions"] == "1 clusters"
    assert rows["T1A01"]["value"] == pytest.approx(10.5)
    assert rows["T9A01"]["exceptions"] == "To be analysed on-site"
